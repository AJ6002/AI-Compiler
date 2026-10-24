"""High-throughput Excel reader implementation using python-calamine (Rust-backed).

Per ADR-007:
- Uses python-calamine for memory-efficient, fast cell matrix extraction.
- Yields CellDetail records matching OpenpyxlReader interface.
- Preserves merged cell topologies and scalar data types without arbitrary code evaluation.
"""

from __future__ import annotations

import datetime
import hashlib
import io
from collections.abc import Iterator
from pathlib import Path
from typing import Any

from openpyxl.utils import get_column_letter
from python_calamine import load_workbook

from src.compiler.contracts.enums import DateEpoch
from src.compiler.contracts.structure import (
    CellDetail,
    RegionBoundary,
    SheetStructure,
    WorkbookStructure,
)
from src.compiler.readers.base import ProbeResult


class CalamineReader:
    """Production-grade Excel reader backed by python-calamine."""

    def probe(self, source: bytes | Path) -> ProbeResult:
        """Fast preliminary inspection of a workbook without heavy cell parsing."""
        if isinstance(source, Path):
            wb = load_workbook(source)
        elif isinstance(source, bytes):
            wb = load_workbook(io.BytesIO(source))
        else:
            wb = load_workbook(source)

        sheet_names = list(wb.sheet_names)
        dims: dict[str, tuple[int, int]] = {}
        for name in sheet_names:
            ws = wb.get_sheet_by_name(name)
            dims[name] = (ws.total_height, ws.total_width)

        return ProbeResult(
            sheet_names=sheet_names,
            sheet_count=len(sheet_names),
            date_epoch="1900",
            is_macro_enabled=False,
            dimensions_map=dims,
        )

    def inspect(
        self,
        source: bytes | Path,
        tenant_id: str,
        project_id: str,
        max_preview_rows: int = 50,
        workbook_name: str | None = None,
    ) -> WorkbookStructure:
        """Extract native layout representation including merged cells and preview rows."""
        if isinstance(source, Path):
            raw_bytes = source.read_bytes()
            wb_name = workbook_name or source.name
            wb = load_workbook(source)
        else:
            raw_bytes = source
            wb_name = workbook_name or "in_memory.xlsx"
            wb = load_workbook(io.BytesIO(raw_bytes))

        source_sha256 = hashlib.sha256(raw_bytes).hexdigest()
        sheets: list[SheetStructure] = []

        for idx, sheet_name in enumerate(wb.sheet_names):
            ws = wb.get_sheet_by_name(sheet_name)
            merged_ranges: list[str] = []
            merged_coords: set[tuple[int, int]] = set()

            for (r1, c1), (r2, c2) in ws.merged_cell_ranges:
                # Calamine ranges are 0-based inclusive: ((min_row, min_col), (max_row, max_col))
                coord_range = (
                    f"{get_column_letter(c1 + 1)}{r1 + 1}:{get_column_letter(c2 + 1)}{r2 + 1}"
                )
                merged_ranges.append(coord_range)
                for row_i in range(r1 + 1, r2 + 2):
                    for col_i in range(c1 + 1, c2 + 2):
                        merged_coords.add((row_i, col_i))

            max_row = ws.total_height
            max_col = ws.total_width

            candidate_regions: list[RegionBoundary] = []
            if max_row > 0 and max_col > 0:
                candidate_regions.append(
                    RegionBoundary(
                        min_row=1,
                        max_row=max_row,
                        min_col=1,
                        max_col=max_col,
                        region_type="table",
                    )
                )

            preview_limit = min(max_preview_rows, max_row)
            cells_sample: list[CellDetail] = []
            all_rows = ws.to_python()

            for r_idx in range(1, preview_limit + 1):
                if r_idx - 1 >= len(all_rows):
                    break
                row_data = all_rows[r_idx - 1]
                for c_idx in range(1, max_col + 1):
                    val = row_data[c_idx - 1] if c_idx - 1 < len(row_data) else None
                    is_m = (r_idx, c_idx) in merged_coords
                    if val is None and not is_m:
                        continue

                    data_type = self._infer_data_type(val)
                    cells_sample.append(
                        CellDetail(
                            coordinate=f"{get_column_letter(c_idx)}{r_idx}",
                            row=r_idx,
                            column=c_idx,
                            data_type=data_type,
                            raw_value=val,
                            formula=None,
                            cached_value=None,
                            number_format=None,
                            is_merged=is_m,
                        )
                    )

            sheets.append(
                SheetStructure(
                    name=sheet_name,
                    index=idx,
                    max_row=max_row,
                    max_col=max_col,
                    merged_ranges=merged_ranges,
                    candidate_regions=candidate_regions,
                    cells_sample=cells_sample,
                )
            )

        return WorkbookStructure(
            tenant_id=tenant_id,
            project_id=project_id,
            source_sha256=source_sha256,
            workbook_name=wb_name,
            date_epoch=DateEpoch.EPOCH_1900,
            sheets=sheets,
        )

    def extract_sheet_cells(
        self,
        source: bytes | Path,
        sheet_name: str,
        batch_size: int = 1000,
    ) -> Iterator[list[CellDetail]]:
        """Yield bounded batches of typed cell coordinates for downstream execution."""
        if isinstance(source, Path):
            wb = load_workbook(source)
        elif isinstance(source, bytes):
            wb = load_workbook(io.BytesIO(source))
        else:
            wb = load_workbook(source)

        if sheet_name not in wb.sheet_names:
            raise KeyError(f"Sheet '{sheet_name}' not found in workbook.")

        ws = wb.get_sheet_by_name(sheet_name)
        merged_coords: set[tuple[int, int]] = set()
        for (r1, c1), (r2, c2) in ws.merged_cell_ranges:
            for row_i in range(r1 + 1, r2 + 2):
                for col_i in range(c1 + 1, c2 + 2):
                    merged_coords.add((row_i, col_i))

        all_rows = ws.to_python()
        batch: list[CellDetail] = []

        for r_idx, row in enumerate(all_rows, start=1):
            for c_idx, val in enumerate(row, start=1):
                is_m = (r_idx, c_idx) in merged_coords
                if val is None and not is_m:
                    continue

                data_type = self._infer_data_type(val)
                batch.append(
                    CellDetail(
                        coordinate=f"{get_column_letter(c_idx)}{r_idx}",
                        row=r_idx,
                        column=c_idx,
                        data_type=data_type,
                        raw_value=val,
                        formula=None,
                        cached_value=None,
                        number_format=None,
                        is_merged=is_m,
                    )
                )

                if len(batch) >= batch_size:
                    yield batch
                    batch = []

        if batch:
            yield batch

    @staticmethod
    def _infer_data_type(val: Any) -> str:
        """Infer standard single-character data type conforming to CellDetail schema."""
        if val is None:
            return "z"
        if isinstance(val, bool):
            return "b"
        if isinstance(val, (int, float)):
            return "n"
        if isinstance(val, (datetime.datetime, datetime.date, datetime.time)):
            return "d"
        if isinstance(val, str):
            return "s"
        return "s"
