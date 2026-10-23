"""Deterministic Excel reader implementation using openpyxl.

Per ADR-007 and AGENTS.md:
- Implements dual-pass inspection to isolate raw formula syntax from cached evaluation.
- Preserves merged cell topologies without synthetic data replication.
- Does NOT execute or eval cell formulas or VBA code (VBA loading is disabled).
- Reads elapsed duration formatting into distinct representation.
"""

from __future__ import annotations

import hashlib
import io
from collections.abc import Iterator
from pathlib import Path

import openpyxl

from src.compiler.contracts.enums import DateEpoch
from src.compiler.contracts.structure import (
    CellDetail,
    RegionBoundary,
    SheetStructure,
    WorkbookStructure,
)
from src.compiler.readers.base import ProbeResult
from src.compiler.readers.date_duration import is_duration_format
from src.compiler.readers.formula_handler import inspect_formula_cell


class OpenpyxlReader:
    """Production-grade Excel reader backed by openpyxl."""

    def probe(self, source: bytes | Path) -> ProbeResult:
        """Fast preliminary inspection of a workbook without heavy cell parsing."""
        raw_bytes = source.read_bytes() if isinstance(source, Path) else source
        stream = io.BytesIO(raw_bytes)

        # Use read-only mode for minimal memory and fast header inspection
        wb = openpyxl.load_workbook(stream, read_only=True, data_only=True)
        try:
            sheet_names = list(wb.sheetnames)
            sheet_count = len(sheet_names)
            is_macro = bool(getattr(wb, "vba_archive", None))

            # Detect date epoch
            is_1904 = bool(getattr(wb.properties, "date1904", False))
            date_epoch = "1904" if is_1904 else "1900"

            dims: dict[str, tuple[int, int]] = {}
            for name in sheet_names:
                ws = wb[name]
                dims[name] = (ws.max_row or 0, ws.max_column or 0)

            return ProbeResult(
                sheet_names=sheet_names,
                sheet_count=sheet_count,
                date_epoch=date_epoch,
                is_macro_enabled=is_macro,
                dimensions_map=dims,
            )
        finally:
            wb.close()

    def inspect(
        self,
        source: bytes | Path,
        tenant_id: str,
        project_id: str,
        max_preview_rows: int = 50,
        workbook_name: str | None = None,
    ) -> WorkbookStructure:
        """Extract native layout representation including merged cells and preview rows."""
        if isinstance(source, Path):
            raw_bytes = source.read_bytes()
            wb_name = workbook_name or source.name
        else:
            raw_bytes = source
            wb_name = workbook_name or "in_memory.xlsx"

        source_sha256 = hashlib.sha256(raw_bytes).hexdigest()

        # Dual-pass loading:
        # Pass 1: data_only=False -> extracts formula text and merged ranges
        # Pass 2: data_only=True -> extracts cached evaluations
        wb_formula = openpyxl.load_workbook(
            io.BytesIO(raw_bytes),
            data_only=False,
            read_only=False,
            keep_vba=False,
        )
        wb_cached = openpyxl.load_workbook(
            io.BytesIO(raw_bytes),
            data_only=True,
            read_only=False,
            keep_vba=False,
        )

        try:
            is_1904 = bool(getattr(wb_formula.properties, "date1904", False))
            date_epoch = DateEpoch.EPOCH_1904 if is_1904 else DateEpoch.EPOCH_1900

            sheets: list[SheetStructure] = []

            for idx, sheet_name in enumerate(wb_formula.sheetnames):
                ws_f = wb_formula[sheet_name]
                ws_c = wb_cached[sheet_name]

                # Extract merged cell ranges
                merged_ranges = [str(r) for r in ws_f.merged_cells.ranges]
                merged_coords: set[tuple[int, int]] = set()
                for r in ws_f.merged_cells.ranges:
                    for row in range(r.min_row, r.max_row + 1):
                        for col in range(r.min_col, r.max_col + 1):
                            merged_coords.add((row, col))

                max_row = ws_f.max_row or 0
                max_col = ws_f.max_column or 0

                # Formulate candidate region
                candidate_regions: list[RegionBoundary] = []
                if max_row > 0 and max_col > 0:
                    candidate_regions.append(
                        RegionBoundary(
                            min_row=1,
                            max_row=max_row,
                            min_col=1,
                            max_col=max_col,
                            region_type="table",
                        )
                    )

                # Sample cell details for inspection preview
                preview_limit = min(max_preview_rows, max_row)
                cells_sample: list[CellDetail] = []

                for r in range(1, preview_limit + 1):
                    for c in range(1, max_col + 1):
                        cell_f = ws_f.cell(row=r, column=c)
                        cell_c = ws_c.cell(row=r, column=c)

                        val_f = cell_f.value
                        val_c = cell_c.value
                        is_m = (r, c) in merged_coords

                        # Keep cells that have content or are part of merged headers
                        if val_f is None and val_c is None and not is_m:
                            continue

                        inspection = inspect_formula_cell(val_f, val_c)
                        if inspection.has_formula:
                            data_type = "f"
                            raw_val = val_f
                            formula_expr = inspection.formula_text
                            cached_val = inspection.cached_value
                        else:
                            raw_val = val_f
                            formula_expr = None
                            cached_val = None
                            if cell_f.data_type == "n" and is_duration_format(cell_f.number_format):
                                data_type = "duration"
                            else:
                                data_type = cell_f.data_type or "z"

                        cells_sample.append(
                            CellDetail(
                                coordinate=cell_f.coordinate,
                                row=r,
                                column=c,
                                data_type=data_type,
                                raw_value=raw_val,
                                formula=formula_expr,
                                cached_value=cached_val,
                                number_format=cell_f.number_format,
                                is_merged=is_m,
                            )
                        )

                sheets.append(
                    SheetStructure(
                        name=sheet_name,
                        index=idx,
                        max_row=max_row,
                        max_col=max_col,
                        merged_ranges=merged_ranges,
                        candidate_regions=candidate_regions,
                        cells_sample=cells_sample,
                    )
                )

            return WorkbookStructure(
                tenant_id=tenant_id,
                project_id=project_id,
                source_sha256=source_sha256,
                workbook_name=wb_name,
                date_epoch=date_epoch,
                sheets=sheets,
            )
        finally:
            wb_formula.close()
            wb_cached.close()

    def extract_sheet_cells(
        self,
        source: bytes | Path,
        sheet_name: str,
        batch_size: int = 1000,
    ) -> Iterator[list[CellDetail]]:
        """Yield bounded batches of typed cell coordinates for downstream execution."""
        raw_bytes = source.read_bytes() if isinstance(source, Path) else source

        wb_formula = openpyxl.load_workbook(
            io.BytesIO(raw_bytes),
            data_only=False,
            read_only=False,
            keep_vba=False,
        )
        wb_cached = openpyxl.load_workbook(
            io.BytesIO(raw_bytes),
            data_only=True,
            read_only=False,
            keep_vba=False,
        )

        try:
            if sheet_name not in wb_formula.sheetnames:
                raise KeyError(f"Sheet '{sheet_name}' not found in workbook.")

            ws_f = wb_formula[sheet_name]
            ws_c = wb_cached[sheet_name]

            merged_coords: set[tuple[int, int]] = set()
            for r in ws_f.merged_cells.ranges:
                for row_idx in range(r.min_row, r.max_row + 1):
                    for col_idx in range(r.min_col, r.max_col + 1):
                        merged_coords.add((row_idx, col_idx))

            batch: list[CellDetail] = []
            for row in ws_f.iter_rows(values_only=False):
                for cell_f in row:
                    if cell_f is None:
                        continue
                    r = cell_f.row
                    c = cell_f.column
                    cell_c = ws_c.cell(row=r, column=c)

                    val_f = cell_f.value
                    val_c = cell_c.value
                    is_m = (r, c) in merged_coords

                    if val_f is None and val_c is None and not is_m:
                        continue

                    inspection = inspect_formula_cell(val_f, val_c)
                    if inspection.has_formula:
                        data_type = "f"
                        raw_val = val_f
                        formula_expr = inspection.formula_text
                        cached_val = inspection.cached_value
                    else:
                        raw_val = val_f
                        formula_expr = None
                        cached_val = None
                        if cell_f.data_type == "n" and is_duration_format(cell_f.number_format):
                            data_type = "duration"
                        else:
                            data_type = cell_f.data_type or "z"

                    batch.append(
                        CellDetail(
                            coordinate=cell_f.coordinate,
                            row=r,
                            column=c,
                            data_type=data_type,
                            raw_value=raw_val,
                            formula=formula_expr,
                            cached_value=cached_val,
                            number_format=cell_f.number_format,
                            is_merged=is_m,
                        )
                    )

                    if len(batch) >= batch_size:
                        yield batch
                        batch = []

            if batch:
                yield batch
        finally:
            wb_formula.close()
            wb_cached.close()
