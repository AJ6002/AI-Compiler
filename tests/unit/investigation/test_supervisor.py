"""Unit tests for SandboxSupervisor using mocked Docker clients."""

from __future__ import annotations

import io
import json
from unittest.mock import MagicMock

import openpyxl

from src.compiler.contracts.enums import JobResultState
from src.compiler.contracts.validation import PlanValidationReport
from src.compiler.investigation.container_spec import HardenedContainerConfig
from src.compiler.investigation.supervisor import SandboxSupervisor
from src.compiler.validation.plan_validator import PlanValidator


def _sample_workbook_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.cell(row=1, column=1, value="TIMESTAMP")
    ws.cell(row=1, column=2, value="TEMP")
    ws.cell(row=2, column=1, value="2024-01-01T00:00:00Z")
    ws.cell(row=2, column=2, value=42.0)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_container_spec_generates_hardened_flags():
    """Verify HardenedContainerConfig enforces non-root, read-only, and cgroup limits."""
    spec = HardenedContainerConfig(image="aiconnex-sandbox:latest")
    kwargs = spec.to_docker_kwargs(input_dir="/host/in", output_dir="/host/out")

    assert kwargs["network_mode"] == "none"
    assert kwargs["user"] == "1000:1000"
    assert kwargs["read_only"] is True
    assert kwargs["cap_drop"] == ["ALL"]
    assert kwargs["security_opt"] == ["no-new-privileges:true"]
    assert kwargs["pids_limit"] == 64
    assert kwargs["mem_limit"] == "512m"
    assert kwargs["memswap_limit"] == "512m"
    assert kwargs["nano_cpus"] == 1_000_000_000
    assert kwargs["tmpfs"]["/work"] == "size=256m"
    assert kwargs["volumes"]["/host/in"]["mode"] == "ro"
    assert kwargs["volumes"]["/host/out"]["mode"] == "rw"


def test_supervisor_handles_timeout():
    """Supervisor kills container on timeout and emits SANDBOX_FAILED(reason=TIMEOUT)."""
    mock_container = MagicMock()
    mock_container.id = "test_container_123"
    mock_container.image.id = "sha256:abcd"
    mock_container.status = "running"
    mock_container.logs.return_value = b"Running long calculation..."

    mock_client = MagicMock()
    mock_client.containers.create.return_value = mock_container

    supervisor = SandboxSupervisor(timeout_s=1, docker_client=mock_client)

    wb_bytes = _sample_workbook_bytes()
    outcome = supervisor.investigate(
        brief_data={"file_name": "test.xlsx"},
        workbook_bytes=wb_bytes,
    )

    assert outcome.status == "SANDBOX_FAILED"
    assert outcome.failure_reason == "TIMEOUT"
    assert outcome.metadata.timed_out is True
    assert outcome.metadata.killed is True
    mock_container.stop.assert_called_with(timeout=5)
    mock_container.remove.assert_called_with(force=True)


def test_supervisor_handles_missing_output():
    """Supervisor marks NO_OUTPUT when container cleanly finishes but emits no result.json."""
    mock_container = MagicMock()
    mock_container.id = "test_container_456"
    mock_container.status = "exited"
    mock_container.wait.return_value = {"StatusCode": 0}
    mock_container.logs.return_value = b""

    mock_client = MagicMock()
    mock_client.containers.create.return_value = mock_container

    supervisor = SandboxSupervisor(timeout_s=10, docker_client=mock_client)

    wb_bytes = _sample_workbook_bytes()
    outcome = supervisor.investigate(
        brief_data={"file_name": "test.xlsx"},
        workbook_bytes=wb_bytes,
        agent_script="print('No output written')",
    )

    assert outcome.status == "SANDBOX_FAILED"
    assert outcome.failure_reason == "NO_OUTPUT"
    mock_container.remove.assert_called_with(force=True)


def test_supervisor_handles_contract_violation():
    """Supervisor marks CONTRACT_VIOLATION when result.json is invalid JSON or violates schema."""
    mock_container = MagicMock()
    mock_container.id = "test_container_789"
    mock_container.status = "exited"
    mock_container.wait.return_value = {"StatusCode": 0}

    mock_client = MagicMock()
    mock_client.containers.create.return_value = mock_container

    supervisor = SandboxSupervisor(timeout_s=10, docker_client=mock_client)

    # Script writing malformed json
    script = "with open('/output/result.json', 'w') as f:\n  f.write('not-json')\n"

    # We mock container execution to write to output dir
    def fake_start():
        call_kwargs = mock_client.containers.create.call_args[1]
        for host_path, bind_info in call_kwargs["volumes"].items():
            if bind_info["bind"] == "/output":
                with open(f"{host_path}/result.json", "w", encoding="utf-8") as f:
                    f.write("not-valid-json")

    mock_container.start.side_effect = fake_start

    wb_bytes = _sample_workbook_bytes()
    outcome = supervisor.investigate(
        brief_data={"file_name": "test.xlsx"},
        workbook_bytes=wb_bytes,
        agent_script=script,
    )

    assert outcome.status == "SANDBOX_FAILED"
    assert outcome.failure_reason == "CONTRACT_VIOLATION"


def test_supervisor_handles_review_required_without_calling_validator():
    """When container returns REVIEW_REQUIRED, supervisor returns it without invoking validator."""
    mock_container = MagicMock()
    mock_container.id = "test_review_1"
    mock_container.status = "exited"
    mock_container.wait.return_value = {"StatusCode": 0}

    mock_client = MagicMock()
    mock_client.containers.create.return_value = mock_container

    mock_validator = MagicMock(spec=PlanValidator)
    supervisor = SandboxSupervisor(
        timeout_s=10,
        docker_client=mock_client,
        plan_validator=mock_validator,
    )

    canned = {
        "schema_version": "1.0.0",
        "tenant_id": "tenant_1",
        "project_id": "proj_1",
        "agent_id": "agent:test",
        "agent_version": "1.0.0",
        "status": "REVIEW_REQUIRED",
        "parse_plan": None,
        "unresolved": [
            {
                "tenant_id": "tenant_1",
                "project_id": "proj_1",
                "issue": "Cannot locate headers",
                "source_file": "sample.xlsx",
                "sheet": "Sheet1",
                "cell_range": "A1:E5",
                "required_clarification": "Define header row",
            }
        ],
        "alternate_interpretations_considered": [],
        "evidence_references": {},
    }

    def fake_start():
        call_kwargs = mock_client.containers.create.call_args[1]
        for host_path, bind_info in call_kwargs["volumes"].items():
            if bind_info["bind"] == "/output":
                with open(f"{host_path}/result.json", "w", encoding="utf-8") as f:
                    json.dump(canned, f)

    mock_container.start.side_effect = fake_start

    wb_bytes = _sample_workbook_bytes()
    outcome = supervisor.investigate(
        brief_data={"file_name": "sample.xlsx"},
        workbook_bytes=wb_bytes,
    )

    assert outcome.status == "REVIEW_REQUIRED"
    assert len(outcome.unresolved) == 1
    assert outcome.unresolved[0].issue == "Cannot locate headers"
    mock_validator.validate.assert_not_called()


def test_supervisor_validates_plan_proposed():
    """When container returns PLAN_PROPOSED, supervisor validates through PlanValidator."""
    mock_container = MagicMock()
    mock_container.id = "test_plan_prop"
    mock_container.status = "exited"
    mock_container.wait.return_value = {"StatusCode": 0}

    mock_client = MagicMock()
    mock_client.containers.create.return_value = mock_container

    mock_validator = MagicMock(spec=PlanValidator)
    mock_validator.validate.return_value = PlanValidationReport(
        plan_id="plan_sample_1",
        status=JobResultState.ACCEPTED,
        schema_passed=True,
        semantics_passed=True,
        evidence_passed=True,
    )

    supervisor = SandboxSupervisor(
        timeout_s=10,
        docker_client=mock_client,
        plan_validator=mock_validator,
    )

    canned = {
        "schema_version": "1.0.0",
        "tenant_id": "tenant_1",
        "project_id": "proj_1",
        "agent_id": "agent:test",
        "agent_version": "1.0.0",
        "status": "PLAN_PROPOSED",
        "parse_plan": {
            "schema_version": "1.0.0",
            "tenant_id": "tenant_1",
            "project_id": "proj_1",
            "plan_id": "plan_sample_1",
            "target_sheet": "Sheet1",
            "operations": [
                {
                    "op": "select_region",
                    "params": {"min_row": 1, "max_row": 2, "min_col": 1, "max_col": 2},
                }
            ],
        },
        "unresolved": [],
        "alternate_interpretations_considered": [],
        "evidence_references": {},
    }

    def fake_start():
        call_kwargs = mock_client.containers.create.call_args[1]
        for host_path, bind_info in call_kwargs["volumes"].items():
            if bind_info["bind"] == "/output":
                with open(f"{host_path}/result.json", "w", encoding="utf-8") as f:
                    json.dump(canned, f)

    mock_container.start.side_effect = fake_start

    wb_bytes = _sample_workbook_bytes()
    outcome = supervisor.investigate(
        brief_data={"file_name": "sample.xlsx"},
        workbook_bytes=wb_bytes,
    )

    assert outcome.status == "PLAN_VALIDATED"
    assert outcome.plan_validation_report is not None
    assert outcome.plan_validation_report.status == JobResultState.ACCEPTED
    mock_validator.validate.assert_called_once()
