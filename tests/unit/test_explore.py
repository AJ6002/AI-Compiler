from pathlib import Path

import openpyxl

from tools.exploration.explore import explore_workbook


def test_explore_workbook(tmp_path: Path):
    # 1. Create a synthetic test workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "IndustrialData"

    # Merge some cells for header band
    ws.merge_cells("A1:B1")
    ws["A1"] = "EQUIPMENT_A"

    # Row 2: column headers
    ws["A2"] = "TIMESTAMP"
    ws["B2"] = "PRESSURE"
    ws["C2"] = "TEMPERATURE"
    ws["D2"] = "PRESSURE"  # Intentional duplicate to test detection

    # Rows 3-5: data
    for r in range(3, 6):
        ws[f"A{r}"] = f"2026-01-0{r}T00:00:00Z"
        ws[f"B{r}"] = 100.5 + r
        ws[f"C{r}"] = 25.0 + r
        ws[f"D{r}"] = 101.0 + r

    file_path = tmp_path / "sample_test.xlsx"
    wb.save(file_path)
    wb.close()

    # 2. Run exploration
    summary = explore_workbook(file_path)

    # 3. Assert structural findings
    assert summary["file_name"] == "sample_test.xlsx"
    assert summary["sheet_count"] == 1
    assert "IndustrialData" in summary["sheets"]

    sheet_info = summary["sheets"]["IndustrialData"]
    assert sheet_info["merged_cell_count"] == 1
    assert "A1:B1" in sheet_info["merged_ranges_sample"]
    assert 2 in sheet_info["candidate_header_rows"]

    # Verify duplicate column name detection in row 2
    assert 2 in sheet_info["duplicate_columns_in_headers"]
    assert "PRESSURE" in sheet_info["duplicate_columns_in_headers"][2]

    # Verify draft recipe
    draft = sheet_info["draft_recipe"]
    assert draft["recipe_version"] == "1.0-draft"
    assert draft["region"]["header_row"] == 2
    assert draft["region"]["data_start_row"] == 3
