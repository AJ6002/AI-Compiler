"""Unit tests for the Three-Level PlanValidator (SCHEMA, SEMANTICS, EVIDENCE)."""

from __future__ import annotations

import io

import openpyxl

from src.compiler.contracts.enums import (
    ErrorCode,
    JobResultState,
    PlanOperationType,
)
from src.compiler.contracts.plans import ParsePlan, PlanOperation
from src.compiler.validation.plan_validator import PlanValidator


def _create_sample_workbook(
    sheet_name: str = "Sheet1",
    include_ambiguous_date: bool = False,
    include_invalid_data: bool = False,
) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = sheet_name

    # Header
    ws.cell(row=1, column=1, value="TIMESTAMP")
    ws.cell(row=1, column=2, value="PRESSURE")
    ws.cell(row=1, column=3, value="TEMPERATURE")

    # Data rows
    if include_ambiguous_date:
        # 01/02/2024 is ambiguous (could be Jan 2 or Feb 1)
        ws.cell(row=2, column=1, value="01/02/2024")
    else:
        ws.cell(row=2, column=1, value="2024-01-02T10:00:00Z")

    if include_invalid_data:
        ws.cell(row=2, column=2, value="not_a_float")
        ws.cell(row=2, column=3, value="also_invalid")
    else:
        ws.cell(row=2, column=2, value=101.3)
        ws.cell(row=2, column=3, value=45.2)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_level_1_schema_failure_short_circuits():
    """Level 1 SCHEMA validation catches missing required parameters and short circuits."""
    validator = PlanValidator()
    # Missing required 'min_row', 'max_row', etc.
    invalid_op = PlanOperation(
        op=PlanOperationType.SELECT_REGION,
        params={"invalid_key": 123},
    )
    plan = ParsePlan(
        tenant_id="tenant_a",
        project_id="proj_1",
        plan_id="plan_schema_fail",
        target_sheet="Sheet1",
        operations=[invalid_op],
    )

    wb_bytes = _create_sample_workbook()
    rep = validator.validate(plan=plan, source=wb_bytes)

    assert rep.status == JobResultState.REVIEW_REQUIRED
    assert rep.schema_passed is False
    assert rep.semantics_passed is False
    assert rep.evidence_passed is False
    assert len(rep.issues) >= 1
    assert rep.issues[0].level == "SCHEMA"
    assert rep.issues[0].code == ErrorCode.ERR_VAL_SCHEMA_VIOLATION


def test_level_2_semantics_sem_004_unpivot_overlap():
    """Level 2 catches overlapping id_vars and value_vars in UNPIVOT (SEM-004)."""
    validator = PlanValidator()
    ops = [
        PlanOperation(
            op=PlanOperationType.SELECT_REGION,
            params={"min_row": 1, "max_row": 2, "min_col": 1, "max_col": 3},
        ),
        PlanOperation(
            op=PlanOperationType.UNPIVOT,
            params={
                "id_vars": ["TIMESTAMP", "PRESSURE"],
                "value_vars": ["PRESSURE", "TEMPERATURE"],  # Overlap on PRESSURE
            },
        ),
    ]
    plan = ParsePlan(
        tenant_id="tenant_a",
        project_id="proj_1",
        plan_id="plan_sem_004",
        target_sheet="Sheet1",
        operations=ops,
    )

    wb_bytes = _create_sample_workbook()
    rep = validator.validate(plan=plan, source=wb_bytes)

    assert rep.status == JobResultState.REVIEW_REQUIRED
    assert rep.schema_passed is True
    assert rep.semantics_passed is False
    assert rep.evidence_passed is False
    assert any(iss.code == ErrorCode.ERR_STRUCT_DUPLICATE_COLUMNS for iss in rep.issues)
    assert any(
        "UNPIVOT id_vars and value_vars must be disjoint" in iss.message for iss in rep.issues
    )


def test_level_2_semantics_sem_003_inverted_boundaries():
    """Level 2 catches inverted boundaries in SELECT_REGION (SEM-003)."""
    validator = PlanValidator()
    ops = [
        PlanOperation(
            op=PlanOperationType.SELECT_REGION,
            params={"min_row": 10, "max_row": 2, "min_col": 1, "max_col": 3},  # min_row > max_row
        ),
    ]
    plan = ParsePlan(
        tenant_id="tenant_a",
        project_id="proj_1",
        plan_id="plan_sem_003",
        target_sheet="Sheet1",
        operations=ops,
    )

    wb_bytes = _create_sample_workbook()
    rep = validator.validate(plan=plan, source=wb_bytes)

    assert rep.status == JobResultState.REVIEW_REQUIRED
    assert rep.schema_passed is True
    assert rep.semantics_passed is False
    assert any(iss.code == ErrorCode.ERR_EXEC_INVALID_RANGE for iss in rep.issues)


def test_level_2_semantics_sem_001_ambiguous_date_tokens():
    """Level 2 flags ambiguous date token e.g. 01/02/2024 without contract specification."""
    validator = PlanValidator()
    ops = [
        PlanOperation(
            op=PlanOperationType.SELECT_REGION,
            params={"min_row": 1, "max_row": 2, "min_col": 1, "max_col": 3},
        ),
        PlanOperation(
            op=PlanOperationType.UNPIVOT,
            params={
                "id_vars": ["TIMESTAMP"],
                "value_vars": ["PRESSURE", "TEMPERATURE"],
            },
        ),
    ]
    plan = ParsePlan(
        tenant_id="tenant_a",
        project_id="proj_1",
        plan_id="plan_sem_001",
        target_sheet="Sheet1",
        operations=ops,
    )

    wb_bytes = _create_sample_workbook(include_ambiguous_date=True)
    rep = validator.validate(plan=plan, source=wb_bytes)

    assert rep.status == JobResultState.REVIEW_REQUIRED
    assert rep.schema_passed is True
    assert rep.semantics_passed is False
    assert any(iss.code == ErrorCode.ERR_STRUCT_AMBIGUOUS_HEADER for iss in rep.issues)
    assert any("Ambiguous date" in iss.message for iss in rep.issues)


def test_level_3_evidence_evd_004_silent_coercion():
    """Level 3 catches cases where all unpivoted values coerced to null (EVD-004)."""
    validator = PlanValidator()
    ops = [
        PlanOperation(
            op=PlanOperationType.SELECT_REGION,
            params={"min_row": 1, "max_row": 2, "min_col": 1, "max_col": 3},
        ),
        PlanOperation(
            op=PlanOperationType.FLATTEN_HEADERS,
            params={"header_rows": [1]},
        ),
        PlanOperation(
            op=PlanOperationType.UNPIVOT,
            params={
                "id_vars": ["TIMESTAMP"],
                "value_vars": ["PRESSURE", "TEMPERATURE"],
            },
        ),
    ]
    plan = ParsePlan(
        tenant_id="tenant_a",
        project_id="proj_1",
        plan_id="plan_evd_004",
        target_sheet="Sheet1",
        operations=ops,
    )

    wb_bytes = _create_sample_workbook(include_invalid_data=True)
    rep = validator.validate(plan=plan, source=wb_bytes)

    assert rep.status == JobResultState.REVIEW_REQUIRED
    assert rep.schema_passed is True
    assert rep.semantics_passed is True
    assert rep.evidence_passed is False
    assert any("Silent coercion failure" in iss.message for iss in rep.issues)


def test_valid_plan_passes_all_three_levels():
    """A valid plan with clean workbook data passes all 3 levels with ACCEPTED."""
    validator = PlanValidator()
    ops = [
        PlanOperation(
            op=PlanOperationType.SELECT_REGION,
            params={"min_row": 1, "max_row": 2, "min_col": 1, "max_col": 3},
        ),
        PlanOperation(
            op=PlanOperationType.FLATTEN_HEADERS,
            params={"header_rows": [1]},
        ),
        PlanOperation(
            op=PlanOperationType.UNPIVOT,
            params={
                "id_vars": ["TIMESTAMP"],
                "value_vars": ["PRESSURE", "TEMPERATURE"],
            },
        ),
    ]
    plan = ParsePlan(
        tenant_id="tenant_a",
        project_id="proj_1",
        plan_id="plan_valid_clean",
        target_sheet="Sheet1",
        operations=ops,
    )

    wb_bytes = _create_sample_workbook()
    rep = validator.validate(plan=plan, source=wb_bytes)

    assert rep.status == JobResultState.ACCEPTED
    assert rep.schema_passed is True
    assert rep.semantics_passed is True
    assert rep.evidence_passed is True
    assert len(rep.issues) == 0
