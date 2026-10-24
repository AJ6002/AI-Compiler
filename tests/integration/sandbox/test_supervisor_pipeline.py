"""Integration test suite executing SandboxSupervisor against Docker with fake agents."""

from __future__ import annotations

import io

import openpyxl
import pytest

import docker
from src.compiler.contracts.enums import JobResultState
from src.compiler.investigation.supervisor import SandboxSupervisor


def _docker_available() -> bool:
    try:
        client = docker.from_env()
        client.ping()
        return True
    except Exception:
        return False


pytestmark = pytest.mark.skipif(not _docker_available(), reason="Docker daemon is not available")


def _sample_workbook_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.cell(row=1, column=1, value="TIMESTAMP")
    ws.cell(row=1, column=2, value="TEMP")
    ws.cell(row=2, column=1, value="2024-01-01T00:00:00Z")
    ws.cell(row=2, column=2, value=42.0)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_supervisor_live_container_valid_plan():
    """End-to-end integration: fake agent produces valid plan, supervisor validates it."""
    supervisor = SandboxSupervisor(image="aiconnex-sandbox:latest", timeout_s=30)

    # Fake agent script that outputs a valid ParsePlan matching workbook structure
    agent_script = """
import json

result = {
  "schema_version": "1.0.0",
  "tenant_id": "tenant_live",
  "project_id": "proj_live",
  "agent_id": "agent:fake_live",
  "agent_version": "1.0.0",
  "status": "PLAN_PROPOSED",
  "parse_plan": {
    "schema_version": "1.0.0",
    "tenant_id": "tenant_live",
    "project_id": "proj_live",
    "plan_id": "plan_live_1",
    "target_sheet": "Sheet",
    "operations": [
      {
        "op": "select_region",
        "params": {"min_row": 1, "max_row": 2, "min_col": 1, "max_col": 2}
      },
      {
        "op": "flatten_headers",
        "params": {"header_rows": [1]}
      },
      {
        "op": "unpivot",
        "params": {
          "id_vars": ["TIMESTAMP"],
          "value_vars": ["TEMP"]
        }
      }
    ]
  },
  "unresolved": [],
  "alternate_interpretations_considered": [],
  "evidence_references": {}
}

with open("/output/result.json", "w") as f:
    json.dump(result, f, indent=2)
"""

    wb_bytes = _sample_workbook_bytes()
    outcome = supervisor.investigate(
        brief_data={"file_name": "live.xlsx"},
        workbook_bytes=wb_bytes,
        workbook_filename="live.xlsx",
        agent_script=agent_script,
    )

    assert outcome.status == "PLAN_VALIDATED"
    assert outcome.failure_reason is None
    assert outcome.investigation_result is not None
    assert outcome.plan_validation_report is not None
    assert outcome.plan_validation_report.status == JobResultState.ACCEPTED
    assert outcome.metadata.timed_out is False
    assert outcome.metadata.exit_code == 0


def test_supervisor_live_container_unresolved_envelope():
    """End-to-end integration: fake agent outputs REVIEW_REQUIRED, supervisor captures issues."""
    supervisor = SandboxSupervisor(image="aiconnex-sandbox:latest", timeout_s=30)

    agent_script = """
import json

result = {
  "schema_version": "1.0.0",
  "tenant_id": "tenant_live",
  "project_id": "proj_live",
  "agent_id": "agent:fake_live",
  "agent_version": "1.0.0",
  "status": "REVIEW_REQUIRED",
  "parse_plan": None,
  "unresolved": [
    {
      "tenant_id": "tenant_live",
      "project_id": "proj_live",
      "issue": "Multi-tiered headers require human disambiguation",
      "source_file": "live.xlsx",
      "sheet": "Sheet",
      "cell_range": "A1:B2",
      "required_clarification": "Which row represents signal units?"
    }
  ],
  "alternate_interpretations_considered": ["Single header row", "Two-tier header row"],
  "evidence_references": {"A1": "Cell contains text 'Timestamp'"}
}

with open("/output/result.json", "w") as f:
    json.dump(result, f, indent=2)
"""

    wb_bytes = _sample_workbook_bytes()
    outcome = supervisor.investigate(
        brief_data={"file_name": "live.xlsx"},
        workbook_bytes=wb_bytes,
        workbook_filename="live.xlsx",
        agent_script=agent_script,
    )

    assert outcome.status == "REVIEW_REQUIRED"
    assert len(outcome.unresolved) == 1
    assert outcome.unresolved[0].issue == "Multi-tiered headers require human disambiguation"
    assert outcome.plan_validation_report is None


def test_supervisor_live_container_cleanup():
    """Verify that after execution, no container with label aiconnex-sandbox is left running."""
    client = docker.from_env()
    containers = client.containers.list(all=True, filters={"label": "aiconnex-sandbox=true"})
    assert len(containers) == 0
